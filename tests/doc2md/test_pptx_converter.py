# Copyright (c) 2026 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import pytest

pptx = pytest.importorskip("pptx")

from pptx.util import Inches  # noqa: E402

from paddleocr._doc2md import convert  # noqa: E402


def _set_strike(run, value):
    run.font._rPr.set("strike", value)


def test_table_cells_are_converted(tmp_path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "apple"
    table.cell(1, 1).text = "3"
    _set_strike(table.cell(1, 1).text_frame.paragraphs[0].runs[0], "dblStrike")
    path = tmp_path / "table.pptx"
    prs.save(str(path))

    markdown = convert(path).markdown

    assert "<th>Name</th>" in markdown
    assert "<td>apple</td>" in markdown
    assert "<td><del>3</del></td>" in markdown


def test_text_frame_strikethrough(tmp_path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.paragraphs[0]
    struck = paragraph.add_run()
    struck.text = "old"
    _set_strike(struck, "sngStrike")
    plain = paragraph.add_run()
    plain.text = " new"
    _set_strike(plain, "noStrike")
    path = tmp_path / "text.pptx"
    prs.save(str(path))

    markdown = convert(path).markdown

    assert markdown == "- ~~old~~ new"
